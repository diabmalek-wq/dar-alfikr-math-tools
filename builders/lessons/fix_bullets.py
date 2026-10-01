import sys
from pptx import Presentation
from lxml import etree
A="http://schemas.openxmlformats.org/drawingml/2006/main"
M="⁣"
for f in sys.argv[1:]:
    prs=Presentation(f); n=0
    for sl in prs.slides:
        for sh in sl.shapes:
            if not sh.has_text_frame: continue
            for p in sh.text_frame.paragraphs:
                pp=p._p.findall("{%s}pPr"%A)
                for extra in pp[1:]: p._p.remove(extra)
                if pp and p._p.index(pp[0])!=0:
                    p._p.remove(pp[0]); p._p.insert(0,pp[0])
                rs=p.runs
                mark=bool(rs) and rs[0].text.startswith(M)
                for r in rs:
                    if M in r.text: r.text=r.text.replace(M,"")
                if mark:
                    pPr=p._p.get_or_add_pPr(); pPr.set("marL","228600"); pPr.set("indent","-228600")
                    for c in list(pPr):
                        if c.tag in ("{%s}buNone"%A,"{%s}buChar"%A,"{%s}buAutoNum"%A): pPr.remove(c)
                    bu=etree.Element("{%s}buChar"%A); bu.set("char","•")
                    pos=len(pPr)
                    for i,c in enumerate(pPr):
                        if c.tag.split('}')[1] in ("tabLst","defRPr","extLst"): pos=i; break
                    pPr.insert(pos,bu); n+=1
    prs.save(f); print(f,"bullets",n)
