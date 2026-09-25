"""Add click-by-click entrance animations to a lesson deck.

pptxgenjs cannot emit animations, so the <p:timing> tree is written straight
into each slide's XML afterwards. Nothing about the layout changes — only the
order in which the existing shapes arrive on screen.

Grouping, which is what makes the order LOGICAL rather than arbitrary:

  chrome      logos, phase tag, time pill, title, subtitle and the footer are
              never animated. They are on screen before the teacher says a word.
  column mode when a slide carries two or more large side-by-side panels (the
              three route cards, a graph beside its READ THE PICTURE panel),
              each column arrives as one click, left to right.
  band mode   otherwise shapes are grouped by their vertical position, so the
              six rows of a worked example arrive one row per click, top to
              bottom — which is the whole point on the method slide.

Usage:  python3 animate_deck.py <deck.pptx> [more.pptx ...]
"""
import copy
import sys
from lxml import etree
from pptx import Presentation

E = 914400.0
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"p": P}

# chrome lives above this line or below the footer line, in inches
TOP_CHROME = 2.20
BOTTOM_CHROME = 6.95

SLIDE_W = 13.333       # inches
BAND_TOL = 0.22        # inches: shapes within this vertical distance are one step
CONTAINER_W = 3.00     # a "panel" is at least this wide
CONTAINER_H = 2.00     # ... and this tall


def geom(sh):
    return (sh.left / E, sh.top / E, sh.width / E, sh.height / E)


def plan(slide):
    """Return a list of click groups; each group is a list of shape ids."""
    live = []
    for sh in slide.shapes:
        if sh.left is None or sh.top is None or sh.width is None or sh.height is None:
            continue
        x, y, w, h = geom(sh)
        if y < TOP_CHROME or y > BOTTOM_CHROME:
            continue
        live.append((sh.shape_id, x, y, w, h))
    if not live:
        return []

    # --- column mode ------------------------------------------------------
    panels = [s for s in live if s[3] >= CONTAINER_W and s[4] >= CONTAINER_H]
    spans = []
    for _id, x, y, w, h in sorted(panels, key=lambda s: s[1]):
        if spans and x < spans[-1][1] - 0.05:      # overlapping horizontally
            spans[-1][1] = max(spans[-1][1], x + w)
        else:
            spans.append([x, x + w])
    if len(spans) >= 2:
        # A bar that spans the whole content width (the orientation decks' "IF YOU
        # FALL BEHIND" strip, a closing rule) belongs to neither column. Assigned
        # by centre it lands in whichever column it is a millimetre nearer, which
        # reads as an accident. Give it its own click, placed by its own height.
        FULLW = 0.62 * SLIDE_W
        wide = [s for s in live if s[3] >= FULLW]
        rest = [s for s in live if s[3] < FULLW]

        groups = [[] for _ in spans]
        for sid, x, y, w, h in rest:
            cx = x + w / 2
            best, bestd = 0, None
            for i, (a, b) in enumerate(spans):
                d = 0.0 if a <= cx <= b else min(abs(cx - a), abs(cx - b))
                if bestd is None or d < bestd:
                    best, bestd = i, d
            groups[best].append((sid, y, x))

        out = []
        for g in groups:
            if g:
                ordered = sorted(g, key=lambda t: (t[1], t[2]))
                out.append((ordered[0][1], ordered[0][2],
                            [sid for sid, _y, _x in ordered]))
        # full-width shapes band together by y, same tolerance as band mode
        wide.sort(key=lambda s: (s[2], s[1]))
        cur, cur_y = [], None
        for sid, x, y, w, h in wide:
            if cur_y is not None and abs(y - cur_y) > BAND_TOL:
                out.append((cur_y, 0.0, cur)); cur, cur_y = [], None
            if cur_y is None:
                cur_y = y
            cur.append(sid)
        if cur:
            out.append((cur_y, 0.0, cur))

        return [g for _y, _x, g in sorted(out, key=lambda t: (t[0], t[1]))]

    # --- band mode --------------------------------------------------------
    live.sort(key=lambda s: (s[2], s[1]))
    groups, cur, cur_y = [], [], None
    for sid, x, y, w, h in live:
        if cur_y is None or abs(y - cur_y) <= BAND_TOL:
            cur.append(sid)
            cur_y = y if cur_y is None else cur_y
        else:
            groups.append(cur)
            cur, cur_y = [sid], y
    if cur:
        groups.append(cur)
    return groups


def _el(tag, **attrs):
    e = etree.SubElement.__self__ if False else etree.Element(f"{{{P}}}{tag}")
    for k, v in attrs.items():
        e.set(k, str(v))
    return e


class Ids:
    def __init__(self, start=3):
        self.n = start

    def __call__(self):
        self.n += 1
        return self.n - 1


def effect(sid, nid, first_in_group):
    """One shape's fade-in. The first shape of a click group carries
    clickEffect; the rest ride along with withEffect so a row arrives whole."""
    par = _el("par")
    ctn = _el("cTn", id=nid(), presetID=10, presetClass="entr", presetSubtype=0,
              fill="hold", grpId=0,
              nodeType="clickEffect" if first_in_group else "withEffect")
    par.append(ctn)
    st = _el("stCondLst"); st.append(_el("cond", delay=0)); ctn.append(st)
    ch = _el("childTnLst"); ctn.append(ch)

    setel = _el("set")
    cb = _el("cBhvr")
    c2 = _el("cTn", id=nid(), dur=1, fill="hold")
    s2 = _el("stCondLst"); s2.append(_el("cond", delay=0)); c2.append(s2)
    cb.append(c2)
    tg = _el("tgtEl"); sp = _el("spTgt", spid=sid); tg.append(sp); cb.append(tg)
    an = _el("attrNameLst"); nm = _el("attrName"); nm.text = "style.visibility"
    an.append(nm); cb.append(an)
    setel.append(cb)
    to = _el("to"); sv = _el("strVal", val="visible"); to.append(sv); setel.append(to)
    ch.append(setel)

    ae = _el("animEffect", transition="in", filter="fade")
    cb2 = _el("cBhvr")
    cb2.append(_el("cTn", id=nid(), dur=400))
    tg2 = _el("tgtEl"); tg2.append(_el("spTgt", spid=sid)); cb2.append(tg2)
    ae.append(cb2)
    ch.append(ae)
    return par


def timing(groups):
    nid = Ids()
    t = _el("timing")
    tn = _el("tnLst"); t.append(tn)
    p0 = _el("par"); tn.append(p0)
    root = _el("cTn", id=1, dur="indefinite", restart="never", nodeType="tmRoot")
    p0.append(root)
    rch = _el("childTnLst"); root.append(rch)
    seq = _el("seq", concurrent=1, nextAc="seek"); rch.append(seq)
    main = _el("cTn", id=2, dur="indefinite", nodeType="mainSeq"); seq.append(main)
    mch = _el("childTnLst"); main.append(mch)

    for grp in groups:
        outer = _el("par"); mch.append(outer)
        oc = _el("cTn", id=nid(), fill="hold"); outer.append(oc)
        os_ = _el("stCondLst"); os_.append(_el("cond", delay="indefinite")); oc.append(os_)
        och = _el("childTnLst"); oc.append(och)
        inner = _el("par"); och.append(inner)
        ic = _el("cTn", id=nid(), fill="hold"); inner.append(ic)
        is_ = _el("stCondLst"); is_.append(_el("cond", delay=0)); ic.append(is_)
        ich = _el("childTnLst"); ic.append(ich)
        for i, sid in enumerate(grp):
            ich.append(effect(sid, nid, i == 0))

    prev = _el("prevCondLst")
    c = _el("cond", evt="onPrev", delay=0)
    tgl = _el("tgtEl"); tgl.append(_el("sldTgt")); c.append(tgl)
    prev.append(c); seq.append(prev)
    nxt = _el("nextCondLst")
    c2 = _el("cond", evt="onNext", delay=0)
    tg2 = _el("tgtEl"); tg2.append(_el("sldTgt")); c2.append(tg2)
    nxt.append(c2); seq.append(nxt)
    return t


def animate(path, verbose=True, skip_ends=True):
    """skip_ends leaves the cover and the closing card static — a title card
    that needs five clicks before the lesson can start is an obstacle."""
    prs = Presentation(path)
    last = len(prs.slides._sldIdLst)
    total = 0
    for n, slide in enumerate(prs.slides, 1):
        if skip_ends and (n == 1 or n == last):
            continue
        groups = plan(slide)
        if not groups:
            continue
        sld = slide._element
        for old in sld.findall(f"{{{P}}}timing"):
            sld.remove(old)
        sld.append(timing(groups))
        total += len(groups)
        if verbose:
            print(f"  slide {n:2d}: {len(groups)} click steps "
                  f"({sum(len(g) for g in groups)} shapes)")
    prs.save(path)
    print(f"{path}: {total} click steps total")


if __name__ == "__main__":
    for f in sys.argv[1:]:
        animate(f)
