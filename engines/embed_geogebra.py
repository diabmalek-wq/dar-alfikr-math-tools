"""Put the GeoGebra clip INTO the Lesson 5-4 deck, on the Production — Explore
slide, and clear out the add-in placeholders it was built with.

Mr Thiab chose the embedded-animation route over the PowerPoint add-in, so the
slide must not still say "GEOGEBRA APPLET / EMBED AREA", and it certainly must
not still carry the literal link geogebra.org/m/PLACEHOLDER — that would go up
on a projector in front of a class.
"""
import subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

DECK = "Gr11_T5_L4_Solving_Radical_Equations.pptx"
CLIP = "GeoGebra_Gr11_5-4_Extraneous.mp4"
POSTER = "geogebra_poster.png"
SLIDE = 12                       # Production — Explore
TEAL = RGBColor(0x17, 0xA1, 0x99)
MUTED = RGBColor(0x5C, 0x6E, 0x6C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# a still that shows the first picture, so the slide does not sit on black
subprocess.run(["ffmpeg", "-y", "-i", CLIP, "-ss", "11.0", "-frames:v", "1", POSTER],
               check=True, capture_output=True)

prs = Presentation(DECK)
s = prs.slides[SLIDE]

DROP = ("GEOGEBRA APPLET", "Insert → Add-ins", "Scan to open",
        "geogebra.org/m/PLACEHOLDER")
removed = []
for sh in list(s.shapes):
    txt = sh.text_frame.text.strip() if sh.has_text_frame else ""
    if any(d in txt for d in DROP):
        sh._element.getparent().remove(sh._element)
        removed.append(txt[:40])
        continue
    # Deleting a text box leaves its BACKING RECTANGLE behind — the first pass
    # left an empty teal slab where the link used to be. Remove the empty shapes
    # at those coordinates too: the QR square, and the link's bar.
    if not txt:
        for (bx, by, label) in ((8.30, 2.42, "[QR placeholder square]"),
                                (8.30, 4.10, "[empty link bar]")):
            if abs(sh.left - Inches(bx)) < Inches(0.06) \
                    and abs(sh.top - Inches(by)) < Inches(0.06):
                sh._element.getparent().remove(sh._element)
                removed.append(label)
                break

# the clip, sized to the panel the placeholder occupied (16:9 inside 6.93 in)
mv = s.shapes.add_movie(CLIP, Inches(0.45), Inches(2.42), Inches(6.93), Inches(3.90),
                        poster_frame_image=POSTER, mime_type="video/mp4")

cap = s.shapes.add_textbox(Inches(0.45), Inches(6.40), Inches(6.93), Inches(0.40))
p = cap.text_frame.paragraphs[0]
r = p.add_run()
r.text = ("Click the picture to play — 40 seconds, no internet needed.")
r.font.size = Pt(13); r.font.italic = True; r.font.name = "Calibri"
r.font.color.rgb = MUTED

# right column: a real, typeable address for student devices
box = s.shapes.add_textbox(Inches(8.30), Inches(2.42), Inches(4.58), Inches(1.45))
tf = box.text_frame; tf.word_wrap = True
p0 = tf.paragraphs[0]
r0 = p0.add_run(); r0.text = "ON YOUR OWN DEVICE"
r0.font.size = Pt(13); r0.font.bold = True; r0.font.name = "Calibri"
r0.font.color.rgb = TEAL
p1 = tf.add_paragraph()
r1 = p1.add_run(); r1.text = "geogebra.org/graphing"
r1.font.size = Pt(20); r1.font.bold = True; r1.font.name = "Calibri"
r1.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = "Type the two equations in yourself, then square them and look again."
r2.font.size = Pt(13); r2.font.name = "Calibri"; r2.font.color.rgb = MUTED

prs.save(DECK)
print("removed:", removed)
print("embedded", CLIP, "on slide", SLIDE + 1)
